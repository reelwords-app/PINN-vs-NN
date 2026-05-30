"""
Converts paper_oil_recovery_DL_benchmark.md to a formatted Word document
(.docx) and inserts all figure images produced by the notebook.

Run:  python build_paper_docx.py
Output: paper_oil_recovery_DL_benchmark.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os, re

# ── Figure catalogue (file → caption) ────────────────────────────────────────
FIGURES = {
    "fig01_target_distribution.png":
        "Figure 1. Target variable (oil recovery factor) distribution: "
        "(a) histogram, (b) box plot, (c) cumulative distribution function.",
    "fig02_feature_distributions.png":
        "Figure 2. Histogram distributions of the 14 reservoir and operational input features.",
    "fig03_correlation_heatmap.png":
        "Figure 3. Pearson correlation heatmap of all input features and the target variable.",
    "fig04_feature_vs_target.png":
        "Figure 4. Scatter plots of each input feature versus oil recovery factor "
        "with Pearson correlation coefficient r annotated.",
    "fig05_dataset_split.png":
        "Figure 5. Dataset partition: 70% training, 15% validation, 15% test (n = 3,306).",
    "fig06a_mlp_arch.png":
        "Figure 6a. MLP architecture block diagram.",
    "fig06b_cnn_arch.png":
        "Figure 6b. CNN-1D architecture block diagram.",
    "fig06c_lstm_arch.png":
        "Figure 6c. LSTM architecture block diagram.",
    "fig06d_tfm_arch.png":
        "Figure 6d. Tabular Transformer architecture block diagram.",
    "fig07_learning_curves.png":
        "Figure 7. Training (left) and validation (right) MSE loss curves for all four architectures.",
    "fig08_metric_bars.png":
        "Figure 8. Test-set performance comparison: RMSE, MAE, R², and MAPE for all four models.",
    "fig09_actual_vs_predicted.png":
        "Figure 9. Scatter plots of actual versus predicted oil recovery factor on the test set.",
    "fig10_residuals.png":
        "Figure 10. Residual distributions (actual − predicted) for all four architectures.",
    "fig11_residual_vs_predicted.png":
        "Figure 11. Residual versus predicted plots (homoscedasticity check).",
    "fig12_abs_error_boxplot.png":
        "Figure 12. Absolute prediction error box plots for all four models on the test set.",
    "fig13_radar_chart.png":
        "Figure 13. Radar chart of normalised performance metrics (outer = better).",
    "fig14_complexity_vs_accuracy.png":
        "Figure 14. Model complexity (trainable parameters) versus test-set RMSE.",
    "fig15_feature_importance.png":
        "Figure 15. Permutation feature importance (mean RMSE increase ± 1 SD) "
        "for the best-performing model, with heatmap.",
    "fig16_error_band.png":
        "Figure 16. Prediction error bands for all four models, "
        "with test samples sorted by actual recovery value.",
}

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

def add_heading(doc, text, level):
    sizes = {1: 14, 2: 13, 3: 12}
    bolds = {1: True, 2: True, 3: True}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    set_font(run, size=sizes.get(level, 12), bold=bolds.get(level, False))
    return p

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0.75)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_font(run)
    return p

def add_figure(doc, path, caption):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(5.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(12)
        r = cap.add_run(caption)
        set_font(r, size=10, italic=True)
    else:
        p = doc.add_paragraph()
        r = p.add_run(f"[Figure not yet generated — run Proxy5.ipynb first]\n{caption}")
        set_font(r, size=10, italic=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def add_table_placeholder(doc, label):
    p = doc.add_paragraph()
    r = p.add_run(f"[{label}]")
    set_font(r, size=11, italic=True, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)

# ── Read markdown ─────────────────────────────────────────────────────────────
with open("paper_oil_recovery_DL_benchmark.md", encoding="utf-8") as f:
    md = f.read()

# ── Build document ────────────────────────────────────────────────────────────
doc = Document()

# Page margins
section = doc.sections[0]
section.page_width  = Inches(8.5)
section.page_height = Inches(11)
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin   = Inches(1.25)
section.right_margin  = Inches(1.25)

# ── Title block ───────────────────────────────────────────────────────────────
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(0)
tp.paragraph_format.space_after  = Pt(8)
tr = tp.add_run(
    "Benchmarking Deep Learning Architectures for Oil Recovery Factor Prediction "
    "in Polymer-Flood Reservoirs: A Unified Comparative Study"
)
set_font(tr, size=16, bold=True)

for meta_line in [
    "Authors: [Author 1]¹, [Author 2]¹, [Author 3]²",
    "Affiliations: ¹[Department of Petroleum Engineering, University Name]  "
    "²[Second Institution]",
    "Corresponding Author: [email@institution.edu]",
    "Target Journal: Journal of Petroleum Science and Engineering / "
    "Geoenergy Science and Engineering",
]:
    mp = doc.add_paragraph()
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = mp.add_run(meta_line)
    set_font(mr, size=10, italic=True)

doc.add_paragraph()

# ── Parse and render markdown sections ───────────────────────────────────────
lines = md.split("\n")
i = 0
in_table = False
table_lines = []

while i < len(lines):
    line = lines[i]

    # Skip the title line (already rendered)
    if line.startswith("# Benchmarking Deep"):
        i += 1
        continue

    # H2
    if line.startswith("## "):
        text = line[3:].strip()
        add_heading(doc, text, 1)
        i += 1
        continue

    # H3
    if line.startswith("### "):
        text = line[4:].strip()
        add_heading(doc, text, 2)
        i += 1
        continue

    # H4
    if line.startswith("#### "):
        text = line[5:].strip()
        add_heading(doc, text, 3)
        i += 1
        continue

    # Horizontal rule — skip
    if line.strip() in ("---", "***", "___"):
        i += 1
        continue

    # Bold metadata lines like **Authors:** — render as body
    if line.startswith("**") and line.endswith("**"):
        p = doc.add_paragraph()
        r = p.add_run(line.strip("*"))
        set_font(r, bold=True, size=11)
        i += 1
        continue

    # Table detection
    if line.strip().startswith("|"):
        table_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip().startswith("|"):
            table_lines.append(lines[i])
            i += 1
        # Remove separator row(s)
        rows = [r for r in table_lines if not re.match(r"^\|\s*[-:]+", r)]
        if not rows:
            continue
        # Parse header
        header = [c.strip() for c in rows[0].strip("|").split("|")]
        ncols = len(header)
        t = doc.add_table(rows=len(rows), cols=ncols)
        t.style = "Table Grid"
        for ci, htext in enumerate(header):
            cell = t.cell(0, ci)
            cell.text = htext
            for para in cell.paragraphs:
                for run in para.runs:
                    set_font(run, size=9, bold=True)
        for ri, row_line in enumerate(rows[1:], start=1):
            cells = [c.strip() for c in row_line.strip("|").split("|")]
            for ci, ctext in enumerate(cells[:ncols]):
                cell = t.cell(ri, ci)
                cell.text = ctext
                for para in cell.paragraphs:
                    for run in para.runs:
                        set_font(run, size=9)
        doc.add_paragraph()
        continue

    # Figure placeholder: [Figure X]
    fig_match = re.match(r"\[Figure (\d+[a-d]?)\]", line.strip())
    if fig_match:
        fn = f"fig{int(fig_match.group(1)):02d}" if fig_match.group(1).isdigit() \
             else f"fig{fig_match.group(1)}"
        # Find matching figure
        for fname, cap in FIGURES.items():
            num = re.search(r"fig(\d+[a-d]?)", fname)
            if num and (num.group(1) == fig_match.group(1) or
                        num.group(1).lstrip("0") == fig_match.group(1)):
                add_figure(doc, fname, cap)
                break
        i += 1
        continue

    # Equation lines ($$...$$)
    if line.strip().startswith("$$"):
        eq_text = line.strip().strip("$")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(eq_text)
        set_font(r, name="Courier New", size=11)
        i += 1
        continue

    # Bullet list
    if line.strip().startswith("- ") or line.strip().startswith("* "):
        p = doc.add_paragraph(style="List Bullet")
        text = line.strip().lstrip("-* ").strip()
        # Strip inline bold
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        r = p.add_run(text)
        set_font(r, size=11)
        i += 1
        continue

    # Numbered list
    if re.match(r"^\d+\.\s", line.strip()):
        p = doc.add_paragraph(style="List Number")
        text = re.sub(r"^\d+\.\s", "", line.strip())
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        r = p.add_run(text)
        set_font(r, size=11)
        i += 1
        continue

    # Empty line
    if not line.strip():
        i += 1
        continue

    # Regular paragraph — strip inline markdown
    text = line.strip()
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)   # bold
    text = re.sub(r"\*(.+?)\*",     r"\1", text)   # italic
    text = re.sub(r"`(.+?)`",       r"\1", text)   # code
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text) # links

    if text:
        add_body(doc, text)
    i += 1

# ── Figures appendix ──────────────────────────────────────────────────────────
doc.add_page_break()
add_heading(doc, "Appendix: Figures", 1)

bp = doc.add_paragraph()
br = bp.add_run(
    "The following figures are produced by running Proxy5.ipynb. "
    "All figures are saved as PNG files in the same directory as this document."
)
set_font(br, size=11, italic=True)

for fname, caption in FIGURES.items():
    doc.add_paragraph()
    add_figure(doc, fname, caption)

# ── Save ─────────────────────────────────────────────────────────────────────
out = "paper_oil_recovery_DL_benchmark.docx"
doc.save(out)
print(f"Saved: {out}")
